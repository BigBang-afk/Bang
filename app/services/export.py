"""CSV / Excel / PDF export for trade logs and performance reports."""
import csv
import os

from app.services.analytics import TradeStats

TRADE_EXPORT_COLUMNS = [
    "id", "trade_date", "trade_time", "direction", "entry_price", "stop_loss", "take_profit",
    "exit_price", "lot_size", "risk_percent", "rr_ratio", "commission", "swap", "spread",
    "profit_loss", "status", "session", "notes",
]


def export_csv(trades, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=TRADE_EXPORT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for t in trades:
            writer.writerow({k: (dict(t).get(k)) for k in TRADE_EXPORT_COLUMNS})
    return path


def export_excel(trades, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "XAUUSD Trades"
    header_fill = PatternFill(start_color="1A1A1F", end_color="1A1A1F", fill_type="solid")
    header_font = Font(color="D4AF37", bold=True)

    ws.append(TRADE_EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for t in trades:
        row = dict(t)
        ws.append([row.get(c) for c in TRADE_EXPORT_COLUMNS])

    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    wb.save(path)
    return path


def export_pdf_report(path, title, account_name, stats: TradeStats, period_label=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    gold = colors.HexColor("#D4AF37")
    dark = colors.HexColor("#0D0D0F")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], textColor=gold, fontSize=20)
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#9A9AA5"))

    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20 * mm, bottomMargin=15 * mm)
    elements = [
        Paragraph(f"{title} — XAUUSD", title_style),
        Paragraph(f"Account: {account_name}    Period: {period_label}", sub_style),
        Spacer(1, 10),
    ]

    kpi_rows = [
        ["Total Trades", stats.total_trades, "Win Rate", f"{stats.win_rate:.1f}%"],
        ["Net Profit", f"{stats.net_profit:.2f}", "Profit Factor", f"{stats.profit_factor:.2f}"],
        ["Expectancy", f"{stats.expectancy:.2f}", "Average RR", f"{stats.average_rr:.2f}"],
        ["Largest Win", f"{stats.largest_win:.2f}", "Largest Loss", f"{stats.largest_loss:.2f}"],
        ["Max Drawdown", f"{stats.max_drawdown:.2f}", "Max DD %", f"{stats.max_drawdown_percent:.2f}%"],
        ["Longest Win Streak", stats.longest_win_streak, "Longest Loss Streak", stats.longest_loss_streak],
    ]
    kpi_table = Table(kpi_rows, colWidths=[110, 90, 110, 90])
    kpi_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#2A2A32")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F5EF")),
        ("TEXTCOLOR", (0, 0), (-1, -1), dark),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 16))

    trade_header = ["Date", "Dir", "Entry", "Exit", "Lots", "RR", "P/L", "Status"]
    trade_rows = [trade_header]
    for r in stats.closed[-40:]:
        trade_rows.append([
            r.get("trade_date"), r.get("direction"), r.get("entry_price"), r.get("exit_price"),
            r.get("lot_size"), r.get("rr_ratio"), f"{float(r.get('profit_loss') or 0):.2f}", r.get("status"),
        ])
    trades_table = Table(trade_rows, colWidths=[65, 35, 55, 55, 40, 35, 55, 55])
    trades_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#2A2A32")),
        ("BACKGROUND", (0, 0), (-1, 0), dark),
        ("TEXTCOLOR", (0, 0), (-1, 0), gold),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ]))
    elements.append(Paragraph("Recent Closed Trades", styles["Heading3"]))
    elements.append(trades_table)

    doc.build(elements)
    return path
